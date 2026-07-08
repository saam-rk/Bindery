"""Assert-based end-to-end check: sample file of each format -> Document -> EPUB -> validate.
Run: python test_pipeline.py
"""
import tempfile
import zipfile
from pathlib import Path

from backend.ai_formatter import chunk_text
from backend.converters import convert, md_to_document
from backend.epub_builder import build_epub, validate_epub

TMP = Path(tempfile.mkdtemp(prefix="bindery-test-"))

MD = """# Biology Notes

Intro paragraph before any chapter.

# Cell Structure

The cell is the basic unit of life. It has **many** parts.

## Organelles

- Nucleus
- Mitochondria

> Structure determines function.

# Photosynthesis

Plants convert light into chemical energy.
"""

TXT = """PHYSICS SUMMARY

NEWTONS LAWS

An object in motion stays in motion unless acted on
by an external force. This line continues the same
paragraph.

1.2 Momentum

Momentum equals mass times velocity.

- conserved in collisions
- vector quantity

THERMODYNAMICS

Heat flows from hot to cold.
"""

HTML = """<html><head><title>Chem Notes</title><script>alert(1)</script></head>
<body><div><h1>Atoms</h1><p style="color:red">Everything is <b>atoms</b>.</p>
<h1>Bonds</h1><p>Ionic and covalent.</p><span>Loose text.</span></div></body></html>"""


def make_docx(path: Path):
    import docx
    d = docx.Document()
    d.add_heading("History Notes", 0)
    d.add_heading("The Renaissance", 1)
    d.add_paragraph("A period of cultural rebirth in Europe.")
    d.add_heading("Key Figures", 2)
    d.add_paragraph("Leonardo da Vinci", style="List Bullet")
    d.add_paragraph("Michelangelo", style="List Bullet")
    d.add_heading("The Enlightenment", 1)
    p = d.add_paragraph("Reason became ")
    p.add_run("central").bold = True
    p.add_run(" to thought.")
    d.save(str(path))


def make_pdf(path: Path):
    import fitz
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 100), "Calculus Notes", fontsize=24)
    page.insert_text((72, 150), "Derivatives", fontsize=18)
    y = 180
    for line in ["The derivative measures the instantaneous rate of", "change of a function."]:
        page.insert_text((72, y), line, fontsize=11)
        y += 16
    page2 = doc.new_page()
    page2.insert_text((72, 100), "Integrals", fontsize=18)
    page2.insert_text((72, 130), "The integral accumulates area under a curve.", fontsize=11)
    doc.save(str(path))


def check_epub(doc, name):
    out = TMP / f"{name}.epub"
    build_epub(doc, out)
    report = validate_epub(out)
    assert report["ok"], f"{name}: validation failed: {report['errors']}"
    with zipfile.ZipFile(out) as z:
        assert z.infolist()[0].filename == "mimetype"
        assert any(n.endswith("nav.xhtml") for n in z.namelist())
        assert any("cover" in n for n in z.namelist())
    return report


# --- markdown
(TMP / "notes.md").write_text(MD, encoding="utf-8")
doc = convert(TMP / "notes.md", "notes.md")
assert doc.title == "Biology Notes", doc.title
titles = [c.title for c in doc.chapters]
assert "Cell Structure" in titles and "Photosynthesis" in titles, titles
assert "<ul>" in doc.chapters[titles.index("Cell Structure")].html
assert "<blockquote>" in doc.chapters[titles.index("Cell Structure")].html
check_epub(doc, "md")
print("markdown ok:", titles)

# --- txt
(TMP / "notes.txt").write_text(TXT, encoding="utf-8")
doc = convert(TMP / "notes.txt", "notes.txt")
titles = [c.title for c in doc.chapters]
assert any("NEWTON" in t for t in titles), titles
body = "".join(c.html for c in doc.chapters)
assert "<ul>" in body and "<h2>" in body
check_epub(doc, "txt")
print("txt ok:", titles)

# --- html
(TMP / "notes.html").write_text(HTML, encoding="utf-8")
doc = convert(TMP / "notes.html", "notes.html")
assert doc.title == "Chem Notes"
titles = [c.title for c in doc.chapters]
assert "Atoms" in titles and "Bonds" in titles, titles
assert "alert" not in "".join(c.html for c in doc.chapters)
assert "style=" not in "".join(c.html for c in doc.chapters)
check_epub(doc, "html")
print("html ok:", titles)

# --- docx
make_docx(TMP / "notes.docx")
doc = convert(TMP / "notes.docx", "notes.docx")
titles = [c.title for c in doc.chapters]
assert any("Renaissance" in t for t in titles), titles
body = "".join(c.html for c in doc.chapters)
assert "<strong>central</strong>" in body, body[:500]
assert "<ul>" in body
check_epub(doc, "docx")
print("docx ok:", titles)

# --- pdf
make_pdf(TMP / "notes.pdf")
doc = convert(TMP / "notes.pdf", "notes.pdf")
titles = [c.title for c in doc.chapters]
assert any("Derivatives" in t or "Calculus" in t for t in titles), titles
body = "".join(c.html for c in doc.chapters)
assert "instantaneous rate of change" in body, body  # line-merge worked
check_epub(doc, "pdf")
print("pdf ok:", titles)

# --- scanned pdf (image-only pages -> Windows OCR fallback)
from backend import ocr
if ocr.available():
    import fitz
    src = fitz.open()
    page = src.new_page()
    page.insert_text((72, 100), "Solar energy is a renewable resource.", fontsize=16)
    page.insert_text((72, 140), "It comes directly from the sun.", fontsize=16)
    pix = page.get_pixmap(matrix=fitz.Matrix(2, 2))
    scan = fitz.open()
    spage = scan.new_page(width=page.rect.width, height=page.rect.height)
    spage.insert_image(spage.rect, pixmap=pix)
    scan.save(str(TMP / "scan.pdf"))
    doc = convert(TMP / "scan.pdf", "scan.pdf")
    body = "".join(c.html for c in doc.chapters).lower()
    assert "renewable" in body and "sun" in body, body[:500]
    check_epub(doc, "scan")
    print("scanned pdf ocr ok")
else:
    print("scanned pdf ocr SKIPPED (winocr not installed)")

# --- chunking + md round trip (AI path without an API call)
big = "\n\n".join(f"# Part {i}\n\nParagraph {i} text. " * 3 for i in range(40))
chunks = chunk_text(big, target=2000)
assert all(len(c) <= 2600 for c in chunks), max(len(c) for c in chunks)
assert "\n\n".join(chunks).replace("\n", "") == big.replace("\n", "")
rt = md_to_document("# T\n\nHello **world**.\n\n# C2\n\nMore.", "fallback")
assert rt.title == "T" and len(rt.chapters) >= 1
print("chunking + roundtrip ok")

# --- send-to-kindle guard (no network: just the not-configured path)
from backend.kindle_mail import configured, send
from backend.models import ConversionError
assert not configured({}) and not configured({"smtp_user": "a", "kindle_email": "b"})
assert configured({"smtp_user": "a", "smtp_pass": "b", "kindle_email": "c"})
try:
    send(TMP / "md.epub", {})
    raise AssertionError("send() without config should raise")
except ConversionError:
    pass
print("send-to-kindle guard ok")

print("\nALL CHECKS PASSED —", TMP)
